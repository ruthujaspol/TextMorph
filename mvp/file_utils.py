import fitz  # PyMuPDF
from docx import Document
import json

def read_txt(file):
    return file.read().decode("utf-8")

def read_pdf(file):
    text = ""
    pdf = fitz.open(stream=file.read(), filetype="pdf")
    for page in pdf:
        text += page.get_text()
    return text.strip()

def read_docx(file):
    doc = Document(file)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text.strip()

def export_txt(text):
    return text.encode("utf-8")

def export_docx(text):
    doc = Document()
    doc.add_paragraph(text)
    output_path = "TextMorph_Output.docx"
    doc.save(output_path)
    return output_path

def export_json(input_text, output_text, mode):
    data = {"mode": mode, "input": input_text, "output": output_text}
    return json.dumps(data, indent=4)
